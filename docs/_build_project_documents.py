from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "project_documents"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"
BLUE = "2563EB"
LIGHT_BLUE = "EAF2FF"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "64748B"
INK = "1F2937"
GREEN = "166534"
AMBER = "92400E"
RED = "991B1B"


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc, running_title, compact=False):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25 if compact else 1.10

    for style_name, size, before, after, color in (
        ("Heading 1", 16, 18 if compact else 16, 10 if compact else 8, BLUE),
        ("Heading 2", 13, 14 if compact else 12, 7 if compact else 6, BLUE),
        ("Heading 3", 12, 10 if compact else 8, 5 if compact else 4, NAVY),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5 if not compact else 0.375)
        style.paragraph_format.first_line_indent = Inches(-0.25 if not compact else -0.188)
        style.paragraph_format.space_after = Pt(8 if not compact else 4)
        style.paragraph_format.line_spacing = 1.167 if not compact else 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run(running_title)
    set_run_font(hr, size=9, bold=True, color=MID_GRAY)
    set_page_number(section.footer.paragraphs[0])

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = "Prepared as an MSc project documentation artifact."


def add_title_block(doc, kicker, title, subtitle, status_line):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(kicker.upper())
    set_run_font(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_run_font(r, size=26, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitle)
    set_run_font(r, size=14, italic=True, color=MID_GRAY)

    for label, value in (
        ("Prepared by", "[Student Name / Register Number]"),
        ("Programme", "MSc Data Science"),
        ("Project status", status_line),
        ("Document date", "1 September 2026"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, bold=True, color=INK)
        vr = p.add_run(value)
        set_run_font(vr, color=INK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(8)
    note.paragraph_format.left_indent = Inches(0.15)
    note.paragraph_format.right_indent = Inches(0.15)
    note._p.get_or_add_pPr().append(OxmlElement("w:shd"))
    note._p.pPr[-1].set(qn("w:fill"), LIGHT_BLUE)
    r = note.add_run("Human review is required. The system supports academic intervention; it does not make decisions about students.")
    set_run_font(r, bold=True, color=NAVY)
    doc.add_page_break()


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.75)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, heading in enumerate(headers):
        cell = header.cells[idx]
        cell.text = ""
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(heading))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, italic=True, color=MID_GRAY)


def build_full_report():
    doc = Document()
    configure_document(doc, "Student Success Analytics | Project Information and Progress")
    add_title_block(
        doc,
        "MSc Data Science Project",
        "Student Success Analytics and Early Warning System",
        "A Machine Learning-Based Decision Support System for Early Academic Intervention",
        "Research prototype completed; Power BI-ready model prepared",
    )

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This project develops a student-support early warning prototype using anonymized historical higher-education data. "
        "It combines assessment progress, previous attempts, study load, registration timing, and Virtual Learning Environment (VLE) activity available by day 40. "
        "A calibrated machine learning model estimates the probability of a later adverse academic outcome, defined as Fail or Withdrawn. "
        "The probability is converted into Low, Medium, and High support-review tiers. Faculty members must check the record and student context before offering support."
    )
    doc.add_paragraph(
        "The completed work includes data integration, cleaning, leakage checks, exploratory analysis, feature engineering, model comparison, calibration, threshold selection, held-out testing, SHAP explanations, subgroup auditing, intervention rules, Streamlit dashboards, Power BI-ready data assets, automated tests, and responsible-use documentation."
    )

    doc.add_heading("1. Project Introduction", level=1)
    doc.add_paragraph(
        "Educational institutions collect information about student assessments, submissions, learning-platform activity, previous attempts, and final outcomes. "
        "When the number of students is large, it is difficult for faculty to review every record at the right time. This project organizes available early information and highlights records that may need human attention."
    )

    doc.add_heading("2. Problem Statement and Motivation", level=1)
    doc.add_paragraph(
        "Academic difficulty is sometimes recognized only after a missed assessment, failure, or withdrawal. A repeatable method is needed to summarize early indicators, estimate later risk, explain why the model produced an estimate, and support timely academic help."
    )
    doc.add_paragraph(
        "The motivation is to provide more time for tutoring, assessment guidance, study planning, access support, or an adviser conversation. The model is deliberately limited because inaccurate predictions can stigmatize students or distribute attention unfairly."
    )

    doc.add_heading("3. Aim and Objectives", level=1)
    doc.add_paragraph("Main aim: to design and evaluate a responsible machine learning decision-support prototype for early academic intervention.")
    for text in (
        "Create a clean and reproducible day-40 student-module snapshot.",
        "Use only information available by the prediction date and prevent target leakage.",
        "Compare baseline and machine learning classification models.",
        "Prioritize High-tier recall while considering faculty review capacity.",
        "Explain model behaviour using SHAP and visible support indicators.",
        "Audit performance across approved demographic groups.",
        "Provide separate faculty, student, and evaluation dashboard views.",
        "Document privacy, fairness, limitations, and human oversight.",
    ):
        add_bullet(doc, text)

    doc.add_heading("4. Scope and Decision Boundary", level=1)
    add_table(doc, ["In scope", "Out of scope"], [
        ["Historical public-data research prototype", "Automatic academic or disciplinary decisions"],
        ["Faculty review prioritization", "Changing grades, enrolment, scholarships, or services"],
        ["Supportive intervention suggestions", "Mental-health, motivation, or character inference"],
        ["Model and fairness evaluation", "Claiming local deployment validity without new evidence"],
    ], [4680, 4680])

    doc.add_heading("5. Alignment with the Given Data Analytics Guidelines", level=1)
    doc.add_paragraph("The project follows the stages listed in the supplied ‘Steps in a Data Analytics Project’ document.")
    add_table(doc, ["Guideline stage", "How this project addresses it"], [
        ["Problem identification", "Defined the need for early and supportive academic review."],
        ["Requirement analysis", "Specified users, outputs, success metrics, privacy rules, and non-automated use."],
        ["Data collection", "Used anonymized OULAD CSV tables from the Open University dataset."],
        ["Data understanding / EDA", "Checked 27,042 snapshots, distributions, missing values, modules, cohorts, and outcomes."],
        ["Data cleaning", "Removed ineligible records, checked duplicates and ranges, and documented missing scores."],
        ["Data preprocessing", "Integrated seven tables, created day-40 features, encoded categories, imputed within folds, and split chronologically."],
        ["Visualization", "Produced EDA outputs, Streamlit views, and a Power BI-ready workbook/theme."],
        ["Model building", "Compared Dummy Prior, Logistic Regression, Random Forest, and Histogram Gradient Boosting."],
        ["Evaluation", "Reported ROC-AUC, PR-AUC, recall, precision, calibration, confusion matrix, and subgroup metrics."],
        ["Interpretation", "Used SHAP and rule-based support indicators without causal claims."],
        ["Dashboard", "Built Faculty, Student, and Evaluation views; specified five Power BI pages."],
        ["Deployment", "Provided local run scripts; production deployment remains outside the MSc prototype."],
        ["Documentation", "Prepared report, model card, ethics protocol, data dictionary, and reproducible code."],
        ["Presentation", "Dashboard screenshots and Power BI assets are ready for demonstration; final slides remain a presentation-stage task."],
    ], [2200, 7160], font_size=9)

    doc.add_heading("6. Dataset Source and Structure", level=1)
    doc.add_paragraph(
        "The primary dataset is the Open University Learning Analytics Dataset (OULAD), originally published by The Open University. "
        "It is anonymized and is available from the official dataset source; copies may also be distributed through Kaggle. The original OULAD publication should be cited even when a Kaggle mirror is used."
    )
    add_table(doc, ["Source table", "Use in the project"], [
        ["studentInfo.csv", "Student, module, presentation, previous attempts, credits, and final outcome."],
        ["studentRegistration.csv", "Registration and unregistration timing."],
        ["assessments.csv", "Assessment dates, types, and weights."],
        ["studentAssessment.csv", "Submission dates and assessment scores."],
        ["studentVle.csv", "Dated student interactions and click counts."],
        ["vle.csv", "Learning activity metadata."],
        ["courses.csv", "Module presentation information."],
    ], [2600, 6760])
    doc.add_paragraph(
        "The processed unit is one student in one module presentation at day 40. The snapshot contains 27,042 rows, 37 columns, four presentations, seven modules, and no duplicate student-module-presentation keys."
    )

    doc.add_heading("7. Outcome Definition", level=1)
    doc.add_paragraph(
        "The binary training outcome is ‘adverse_outcome’. Fail and Withdrawn after the day-40 cutoff are coded as 1; Pass and Distinction are coded as 0. "
        "Final outcomes are used only as labels and never as prediction inputs. The model output is a probability, which is converted into support-review tiers using thresholds selected on validation data."
    )
    add_table(doc, ["Tier", "Definition", "Meaning"], [
        ["Low", "Probability below 0.3882", "Continue normal monitoring."],
        ["Medium", "Probability from 0.3882 to below 0.5915", "Consider an academic check-in after review."],
        ["High", "Probability at or above 0.5915", "Priority for faculty review; not a factual label."],
    ], [1200, 3000, 5160])

    doc.add_heading("8. Prediction Features", level=1)
    add_table(doc, ["Feature group", "Features used"], [
        ["Previous study and workload", "Number of previous attempts, studied credits, registration date."],
        ["Assessment progress", "Assessments due/submitted, late assessments, submitted weight, mean/minimum score, weighted score, completion rates."],
        ["VLE engagement", "Total clicks, active days, recent 7- and 14-day clicks, previous-week clicks, recent active days, days since last activity, trend, no-activity flag."],
        ["Academic context", "Module code and presentation season."],
    ], [2200, 7160])
    doc.add_paragraph(
        "Gender, age band, disability, highest education, and deprivation-band fields were excluded from the prediction model and retained separately for authorized fairness analysis. OULAD does not contain ordinary classroom attendance, GPA, or backlog fields, so these were not invented or renamed. VLE activity is treated only as an engagement indicator."
    )

    doc.add_heading("9. Data Processing and Quality Controls", level=1)
    for text in (
        "Included students registered by day 40 with a supported final outcome.",
        "Excluded withdrawals already known on or before day 40.",
        "Aggregated only assessments due and VLE events observed by the cutoff.",
        "Checked duplicate keys, ranges, missingness, and impossible completion rates.",
        "Handled missing assessment scores within training folds to avoid information leakage.",
        "Used chronological cohorts rather than a random split.",
    ):
        add_number(doc, text)
    doc.add_paragraph(
        "Important profile values: 11,663 of 27,042 records had a later adverse outcome (43.1%). The median day-40 VLE activity was 284 clicks across 17 active days. Missing assessment averages occurred in 16.42% of rows, mainly where no scored assessment was yet available."
    )

    doc.add_heading("10. Modelling Method", level=1)
    doc.add_paragraph(
        "Training presentations were 2013B and 2013J (11,628 rows), validation used 2014B (6,366 rows), and the untouched test cohort was 2014J (9,048 rows). "
        "Model and threshold decisions were made using training and validation data only."
    )
    add_table(doc, ["Model", "Reason for comparison"], [
        ["Dummy Prior", "Uninformed baseline."],
        ["Logistic Regression", "Interpretable linear baseline."],
        ["Random Forest", "Captures nonlinear relationships and interactions."],
        ["Histogram Gradient Boosting", "Efficient nonlinear boosting benchmark."],
    ], [3000, 6360])
    doc.add_paragraph("The calibrated Random Forest was selected because it narrowly achieved the best validation ranking under the project’s High-recall and review-capacity policy.")

    doc.add_heading("11. Test Results", level=1)
    add_table(doc, ["Metric", "Held-out 2014J result", "Interpretation"], [
        ["ROC-AUC", "0.761", "Useful ranking ability, but not perfect separation."],
        ["Precision-recall AUC", "0.711", "Performance under class imbalance."],
        ["Brier score", "0.193", "Probability error; lower is better."],
        ["High-tier recall", "45.5%", "High tier captured 45.5% of later adverse outcomes."],
        ["High-tier precision", "72.5%", "72.5% of High-tier records later had the adverse outcome."],
        ["High review rate", "24.9%", "2,253 of 9,048 test records were High."],
        ["Medium + High recall", "74.5%", "Broader review tiers captured 74.5% of adverse outcomes."],
    ], [2300, 2200, 4860], font_size=9)
    doc.add_paragraph(
        "These are historical test results, not promises about future students. High-tier recall below 50% means many students with later adverse outcomes were not placed in High, so the system must not replace normal academic support processes."
    )

    doc.add_heading("12. Explainability and Intervention Support", level=1)
    doc.add_paragraph(
        "SHAP was used to show how model features contributed to predictions globally and for a selected test sample. SHAP values explain the model calculation; they do not prove why a student experienced an outcome."
    )
    add_table(doc, ["Visible indicator", "Faculty-reviewed support option"], [
        ["No recent VLE activity", "Check access, engagement, and study planning."],
        ["Incomplete due assessment", "Verify the record and discuss deadline or academic support."],
        ["Low available assessment score", "Offer feedback, tutoring, or study-skills resources."],
        ["Recent engagement decline", "Confirm the pattern and consider a supportive check-in."],
        ["Previous attempt", "Consider an academic-adviser review."],
    ], [3000, 6360])

    doc.add_heading("13. Dashboard and Decision-Support Interface", level=1)
    doc.add_paragraph(
        "The Streamlit prototype contains Faculty, Student, and Evaluation views. The Faculty view shows the review queue, support tier, probability, indicators, support options, and a permanent human-review warning. "
        "The Student view avoids risk labels and focuses on understandable learning records and support. The Evaluation view reports model performance, calibration, SHAP factors, and subgroup audits."
    )
    faculty_image = ROOT / "artifacts" / "dashboard_faculty_final.png"
    if faculty_image.exists():
        doc.add_picture(str(faculty_image), width=Inches(6.45))
        add_caption(doc, "Figure 1. Faculty decision-support dashboard prototype.")
    evaluation_image = ROOT / "artifacts" / "dashboard_evaluation.png"
    if evaluation_image.exists():
        doc.add_picture(str(evaluation_image), width=Inches(6.45))
        add_caption(doc, "Figure 2. Research evaluation dashboard prototype.")
    doc.add_paragraph(
        "A Power BI-ready Excel data model, DAX measure list, visual theme, and five-page dashboard specification have also been prepared. Native .pbix assembly remains to be completed on a computer with Power BI Desktop."
    )

    doc.add_heading("14. Privacy, Fairness, and Responsible Use", level=1)
    for text in (
        "Use only the minimum necessary data and keep direct identifiers outside modelling data.",
        "Restrict student-level access and maintain review and access logs.",
        "Allow students to correct inaccurate source records.",
        "Do not use the model for discipline, grading, admissions, scholarships, or denial of services.",
        "Review group recall, precision, selection rate, and false-positive rate before any local pilot.",
        "Communicate uncertainty and keep faculty responsible for the final decision.",
    ):
        add_bullet(doc, text)
    doc.add_paragraph(
        "The subgroup audit showed similar High-tier recall by gender but different false-positive rates, which requires investigation rather than a simple declaration of fairness. Demographic differences are monitoring signals, not explanations of student behaviour."
    )

    doc.add_heading("15. Testing and Reproducibility", level=1)
    doc.add_paragraph(
        "The repository contains run scripts for data preparation, exploratory analysis, model training, support outputs, SHAP explanations, and dashboard launch. Automated tests cover schema checks, label creation, risk tiers, feature construction, and intervention rules. The complete test suite passed 13 tests during the latest verified run."
    )

    doc.add_heading("16. Current Project Update", level=1)
    add_table(doc, ["Work item", "Status", "Evidence / output"], [
        ["Problem, scope, objectives", "Completed", "Project charter and final report."],
        ["Dataset selection and provenance", "Completed", "OULAD selection document and raw source tables."],
        ["Data cleaning and preprocessing", "Completed", "27,042-row day-40 snapshot and profile."],
        ["EDA", "Completed", "Summary JSON, CSV tables, and HTML report."],
        ["Model comparison and calibration", "Completed", "Selected calibrated Random Forest and evaluation files."],
        ["SHAP and fairness audit", "Completed", "Global/local SHAP and subgroup audit artifacts."],
        ["Intervention rules", "Completed", "Transparent support-indicator rules."],
        ["Streamlit dashboard", "Completed", "Faculty, Student, and Evaluation views."],
        ["Power BI preparation", "Completed", "Import-ready workbook, DAX measures, theme, and build guide."],
        ["Native Power BI .pbix", "Pending", "Requires Power BI Desktop to assemble and save the report."],
        ["Final dissertation and presentation", "In progress", "Content exists; institutional formatting and slides remain."],
        ["Institutional deployment", "Not started / out of scope", "Requires approval, local data, security, and prospective validation."],
    ], [3000, 1800, 4560], font_size=8.8)

    doc.add_heading("17. Main Deliverables", level=1)
    for text in (
        "Processed OULAD day-40 dataset and feature pipeline.",
        "EDA report, trained model bundle, metrics, review queue, SHAP explanations, and fairness audit.",
        "Faculty, Student, and Evaluation Streamlit dashboard.",
        "Power BI-ready workbook, DAX measures, dashboard plan, and visual theme.",
        "Project charter, dataset selection, data dictionary, model card, ethics protocol, and final project report.",
        "Automated tests and one-command PowerShell run scripts.",
    ):
        add_bullet(doc, text)

    doc.add_heading("18. Limitations and Future Work", level=1)
    doc.add_paragraph(
        "OULAD represents distance-learning modules from one institution. VLE activity does not measure offline study or classroom attendance. The model cannot observe personal circumstances, and Fail and Withdrawn may have different causes. "
        "The selected model only narrowly exceeded Logistic Regression during validation. Public-data results do not establish performance at another university."
    )
    for text in (
        "Compare prediction cutoffs at days 20, 40, and 60.",
        "Model failure and withdrawal separately.",
        "Evaluate simpler models for possible local use.",
        "Complete the native Power BI report and add role-based security for demonstration.",
        "Conduct approved local validation and user testing with students and faculty.",
        "Study whether interventions are timely, acceptable, and beneficial.",
    ):
        add_bullet(doc, text)

    doc.add_heading("19. Conclusion", level=1)
    doc.add_paragraph(
        "The project now demonstrates the complete technical and responsible-design workflow required by the supplied data analytics guidelines. It converts public historical data into a leakage-aware early-warning prototype, evaluates multiple models, explains outputs, audits groups, and presents support options through dashboards. "
        "Its correct role is to help faculty organize supportive review while keeping people responsible for every final decision."
    )

    doc.add_heading("References", level=1)
    refs = [
        "Kuzilek, J., Hlosta, M., & Zdrahal, Z. (2017). Open University Learning Analytics Dataset. Scientific Data, 4, 170171. https://doi.org/10.1038/sdata.2017.171",
        "Lundberg, S. M., & Lee, S.-I. (2017). A Unified Approach to Interpreting Model Predictions. Advances in Neural Information Processing Systems, 30.",
        "Scikit-learn documentation. https://scikit-learn.org/",
        "SHAP documentation. https://shap.readthedocs.io/",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)

    out = OUT / "01_Project_Information_and_Progress_Report.docx"
    doc.save(out)
    return out


def build_book_record():
    doc = Document()
    configure_document(doc, "Student Success Analytics | Project Book Record", compact=True)
    add_title_block(
        doc,
        "MSc Data Science Project Book",
        "Abstract, Dataset, Features and Daily Work Record",
        "Student Success Analytics and Early Warning System",
        "Easy-language record of completed project work",
    )

    doc.add_heading("1. Project Abstract", level=1)
    doc.add_paragraph(
        "This project is called Student Success Analytics and Early Warning System. The main purpose is to help faculty members identify students who may need academic support at an early stage. The system studies information available during the course, such as assessment submission, marks, previous attempts, study credits, and activity in the online learning system."
    )
    doc.add_paragraph(
        "The project uses machine learning to estimate the chance of a later Fail or Withdrawn result. It places each record in Low, Medium, or High support risk. These levels are only used to arrange faculty review. They are not final decisions and they do not mean that a student will definitely fail. Faculty members must check the record and talk to the student before choosing any support."
    )
    doc.add_paragraph(
        "Different models were tested, and a calibrated Random Forest model was selected using validation data. On the separate test data, the model obtained ROC-AUC 0.761 and precision-recall AUC 0.711. The High group found 45.5% of the later adverse outcomes, and 72.5% of High-group records later had an adverse outcome. Medium and High together found 74.5% of adverse outcomes. The project also includes SHAP explanations, fairness checks, support suggestions, and dashboards."
    )

    doc.add_heading("2. Where We Got the Data", level=1)
    doc.add_paragraph(
        "We used the Open University Learning Analytics Dataset, commonly called OULAD. The original data was published by The Open University and is anonymized. It can be downloaded from the official OULAD source, and copies are also available on Kaggle. If a Kaggle copy is used, the original Open University source must still be mentioned."
    )
    doc.add_paragraph("The dataset contains seven connected CSV files:")
    for text in (
        "studentInfo.csv - student, module, previous attempt, credits, and final result information.",
        "studentRegistration.csv - registration and withdrawal dates.",
        "assessments.csv - assessment date, type, and weight.",
        "studentAssessment.csv - student submission date and score.",
        "studentVle.csv - dated online-learning clicks for each student.",
        "vle.csv - information about online-learning activities.",
        "courses.csv - module presentation information.",
    ):
        add_bullet(doc, text)
    doc.add_paragraph(
        "We joined these files and created one row for one student in one module presentation. We used only the information available up to day 40. The prepared data has 27,042 rows."
    )

    doc.add_heading("3. Output We Are Predicting", level=1)
    doc.add_paragraph(
        "The model predicts a later adverse academic outcome. Fail and Withdrawn are placed in the adverse group. Pass and Distinction are placed in the successful group. Final results are used only to train and test the model; they are not used as input features."
    )
    add_table(doc, ["Support level", "Simple meaning"], [
        ["Low", "The record has a lower model-estimated need for review."],
        ["Medium", "The record may need a normal academic check-in."],
        ["High", "The record should be reviewed first by faculty."],
    ], [1800, 7560])

    doc.add_heading("4. Features We Are Taking", level=1)
    add_table(doc, ["Feature group", "Features used in simple words"], [
        ["Previous study", "Number of previous attempts and date of registration."],
        ["Study load", "Number of credits studied."],
        ["Assessment work", "Assessments due, submitted, late, and completed by day 40."],
        ["Assessment marks", "Average score, minimum score, weighted score, and submitted assessment weight."],
        ["Online activity", "Total clicks, active days, clicks in the last 7 and 14 days, and recent active days."],
        ["Activity change", "Previous-week clicks, change in clicks, and days since the last activity."],
        ["Course details", "Module code and presentation season."],
    ], [2200, 7160], font_size=9.2)
    doc.add_paragraph(
        "Gender, age band, disability, highest education, and deprivation band are not used to make the prediction. They are kept separately only to check whether the model shows different error patterns for different groups."
    )
    doc.add_paragraph(
        "Important note: OULAD does not have normal classroom attendance, GPA, or backlog columns. We did not create false values for these fields. Online clicks are treated as online engagement, not as classroom attendance."
    )

    doc.add_heading("5. Simple Project Method", level=1)
    method_steps = [
        "Understand the student-support problem and decide the project scope.",
        "Collect and study the OULAD CSV files.",
        "Clean the data and join the related tables.",
        "Create day-40 assessment and online-activity features.",
        "Check missing values, duplicates, data types, and distributions.",
        "Split the data by time into training, validation, and testing groups.",
        "Train and compare four classification models.",
        "Select model thresholds using validation data only.",
        "Evaluate the final model on untouched test data.",
        "Use SHAP to explain model contributions.",
        "Check fairness indicators for approved groups.",
        "Create faculty, student, and evaluation dashboards.",
        "Prepare Power BI data, measures, theme, and report-page plan.",
        "Test the code and prepare project documents.",
    ]
    add_table(doc, ["Step", "Activity"], [[i + 1, text] for i, text in enumerate(method_steps)], [900, 8460], font_size=8.8)

    doc.add_heading("6. Day-to-Day Work Record", level=1)
    doc.add_paragraph(
        "The following record is written by work session because exact calendar dates were not recorded in the project files. Dates can be added in the left column of the physical project book."
    )
    daily_entries = [
        ("Day 1", "Topic Selection and Scope", "Discussed possible MSc Data Science project topics and selected Student Success Analytics and Early Warning System. Wrote the basic problem statement, main aim, expected users, and project scope. Confirmed that the system will support faculty review and will not make automatic decisions about students."),
        ("Day 2", "Guideline Study", "Read the given data analytics and data preprocessing guidelines carefully. Matched the project with problem identification, requirement analysis, data collection, cleaning, preprocessing, modelling, evaluation, dashboard development, and documentation. Prepared a clear workflow so that every required stage would be covered."),
        ("Day 3", "Dataset Selection", "Studied different public student-performance datasets available from official sources and Kaggle. Selected the Open University Learning Analytics Dataset because it has dated assessment and online-learning activity records. Recorded that the original source is The Open University, even if a Kaggle mirror is used."),
        ("Day 4", "Project Setup", "Created the Python project folders for source code, raw data, processed data, documents, tests, configuration, dashboards, and model outputs. Prepared the virtual environment and requirements file. Added the first project charter, dataset plan, data dictionary, and responsible-use notes."),
        ("Day 5", "Understanding the Raw Data", "Collected the seven OULAD CSV files and studied the meaning of their columns. Checked how student information, registration, courses, assessments, submissions, and online-learning activity are connected. Identified the keys required to combine student, module, and presentation records correctly."),
        ("Day 6", "Data Integration", "Developed the data-integration process to join the related OULAD tables. Combined student details, registration timing, assessment information, submission records, and VLE activity using the correct identifiers. Created one consistent student-module-presentation record for later feature engineering."),
        ("Day 7", "Data Cleaning and Validation", "Checked the raw and combined data for missing values, duplicate keys, invalid ranges, and inconsistent outcomes. Removed records that were not suitable for the selected prediction setup. Added validation checks so that negative counts, impossible rates, and duplicate student-module-presentation rows could be detected."),
        ("Day 8", "Prediction Cutoff and Leakage Control", "Selected day 40 of the module presentation as the early prediction point. Excluded students whose withdrawal was already known on or before day 40. Ensured that final results and any information recorded after day 40 were not used as model features, preventing target leakage."),
        ("Day 9", "Assessment Feature Engineering", "Created assessment features using only work due by day 40. Calculated assessments due, assessments submitted, late assessments, submitted assessment weight, average score, minimum score, weighted score, and completion rates. Checked that completion rates did not exceed the valid range."),
        ("Day 10", "Online-Activity Features", "Processed the large studentVle.csv file in chunks to control memory use. Created total clicks, active days, last 7-day and 14-day clicks, previous-week clicks, recent active days, days since last activity, activity trend, and a no-activity flag. Treated these values as online engagement, not attendance."),
        ("Day 11", "Exploratory Data Analysis", "Performed exploratory data analysis on the prepared day-40 dataset. Studied the number of records, modules, presentations, outcomes, missing values, numerical summaries, and engagement distributions. Generated reusable CSV, JSON, and HTML EDA outputs for the report and dashboard."),
        ("Day 12", "Target Definition", "Created the binary target called adverse_outcome. Fail and Withdrawn after the cutoff were placed in the adverse group, while Pass and Distinction were placed in the successful group. Documented that this target is an observed historical outcome and not a diagnosis of a student."),
        ("Day 13", "Chronological Data Split", "Divided the dataset by academic presentation instead of using a random split. Used 2013B and 2013J for training, 2014B for model validation, and 2014J as the untouched test cohort. This setup gives a more realistic check of performance on later student records."),
        ("Day 14", "Model Training", "Built four classification baselines: Dummy Prior, Logistic Regression, Random Forest, and Histogram Gradient Boosting. Applied categorical encoding, missing-value handling, and other preprocessing inside the model pipeline. Trained each model using the training presentations without using the final test data."),
        ("Day 15", "Model Comparison and Thresholds", "Compared the models using validation ROC-AUC, average precision, probability error, High-risk recall, High-risk precision, and review rate. Calibrated the model probabilities and selected the Random Forest. Chose Low, Medium, and High thresholds using validation data and assumed faculty review capacity."),
        ("Day 16", "Final Model Evaluation", "Evaluated the selected calibrated Random Forest once on the untouched 2014J test cohort. Recorded ROC-AUC 0.761, precision-recall AUC 0.711, High recall 45.5%, High precision 72.5%, and Medium-plus-High recall 74.5%. Clearly stated that these are historical public-data results."),
        ("Day 17", "SHAP Explainability", "Generated global and sampled student-level SHAP explanations for the selected model. Identified which available assessment, activity, study-load, and module features contributed most to model estimates. Documented that SHAP explains the model calculation and does not prove the cause of a student's outcome."),
        ("Day 18", "Fairness and Ethics Review", "Checked group-level outcome rates, selection rates, recall, precision, and false-positive rates for approved demographic fields. Kept these sensitive fields outside the prediction feature list. Prepared privacy, access-control, correction, retention, human-review, and prohibited-use rules for any future institutional pilot."),
        ("Day 19", "Support and Intervention Rules", "Created transparent support indicators from visible academic records, such as incomplete assessments, low available scores, no recent activity, or declining engagement. Mapped each indicator to possible faculty-reviewed support, including tutoring, feedback, access help, study planning, or an adviser check-in. No action is automatic."),
        ("Day 20", "Streamlit Dashboard", "Built separate Faculty, Student, and Evaluation views in Streamlit. The Faculty view contains the prioritized review queue, indicators, and support options. The Student view avoids risk labels and focuses on learning records and support, while the Evaluation view presents model results and fairness information."),
        ("Day 21", "Testing and Run Scripts", "Added automated tests for schemas, labels, feature engineering, risk tiers, and intervention rules. Corrected the complete PowerShell workflow so that the large prepared dataset can be reused without rebuilding it each time. Verified that all 13 automated tests passed in the latest run."),
        ("Day 22", "Power BI Preparation", "Prepared an Excel data model that can be imported into Power BI. Added fact and dimension tables, DAX measures, a colour theme, relationships, and a five-page dashboard plan. Noted that the native .pbix file must be assembled on a computer where Power BI Desktop is installed."),
        ("Day 23", "Final Documentation and Update", "Collected the completed work into the project report, model card, ethics protocol, dataset notes, Power BI guide, and progress update. Checked that the project follows the provided analytics guidelines. Listed the remaining work: final university formatting, Power BI assembly, presentation slides, and demonstration practice."),
    ]
    for day, title, description in daily_entries:
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(3)
        heading.paragraph_format.keep_with_next = True
        day_run = heading.add_run(f"{day} - ")
        set_run_font(day_run, size=12, bold=True, color=BLUE)
        title_run = heading.add_run(title)
        set_run_font(title_run, size=12, bold=True, color=NAVY)
        entry = doc.add_paragraph(description)
        entry.paragraph_format.space_after = Pt(8)
        entry.paragraph_format.line_spacing = 1.15
        entry.paragraph_format.keep_together = True

    doc.add_heading("7. Main Results", level=1)
    add_table(doc, ["Item", "Result"], [
        ["Prepared dataset", "27,042 day-40 student-module records"],
        ["Training data", "11,628 records"],
        ["Validation data", "6,366 records"],
        ["Test data", "9,048 records"],
        ["Selected model", "Calibrated Random Forest"],
        ["ROC-AUC", "0.761"],
        ["Precision-recall AUC", "0.711"],
        ["High recall", "45.5%"],
        ["High precision", "72.5%"],
        ["Medium + High recall", "74.5%"],
    ], [3100, 6260])

    doc.add_heading("8. Work Completed and Work Remaining", level=1)
    doc.add_paragraph("Completed work:")
    for text in (
        "Dataset collection, cleaning, joining, preprocessing, and EDA.",
        "Model training, comparison, calibration, testing, SHAP, and fairness audit.",
        "Support rules, review queue, Streamlit dashboard, Power BI-ready files, and documentation.",
        "Thirteen automated tests passed in the latest verified run.",
    ):
        add_bullet(doc, text)
    doc.add_paragraph("Remaining work:")
    for text in (
        "Open the prepared model in Power BI Desktop and save the native .pbix dashboard.",
        "Arrange the report in the university’s final dissertation format.",
        "Prepare final presentation slides and demonstration practice.",
        "Use local student data only after formal approval and new validation; this is not part of the current prototype.",
    ):
        add_bullet(doc, text)

    doc.add_heading("9. Conclusion", level=1)
    doc.add_paragraph(
        "The project follows the full data analytics workflow: problem selection, requirement study, data collection, understanding, cleaning, preprocessing, visualization, model building, evaluation, interpretation, dashboard development, testing, and documentation. "
        "The system helps faculty decide which records to review first, but faculty remain responsible for checking the facts and choosing suitable academic support."
    )

    doc.add_heading("Reference", level=1)
    doc.add_paragraph(
        "Kuzilek, J., Hlosta, M., & Zdrahal, Z. (2017). Open University Learning Analytics Dataset. Scientific Data, 4, 170171. https://doi.org/10.1038/sdata.2017.171"
    )

    out = OUT / "02_Project_Book_Expanded_Daily_Record.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    first = build_full_report()
    second = build_book_record()
    print(first)
    print(second)
